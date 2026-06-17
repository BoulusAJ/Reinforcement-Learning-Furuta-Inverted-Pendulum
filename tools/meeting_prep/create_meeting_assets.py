from pathlib import Path
import csv
import math
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "meeting_update_2026-06-17"
PI_RUN = ROOT / "results" / "TD3" / "run_20260616_012625_td3_mathworks_style_wide"
V_RUN = ROOT / "results" / "TD3" / "run_20260616_233822_td3_mathworks_style_voltage"
PI_ANALYSIS = PI_RUN / "analysis" / "pi_current_analytical_active"
V_ANALYSIS = V_RUN / "analysis" / "voltage_analytical_active_2"
WATER_ROOT = Path("C:/Users/abuj/Code/Work/RL/RL-Water-Tank")
WATER_RUN3 = WATER_ROOT / "data" / "run 3 - 20260526_1"
WATER_RUN4 = WATER_ROOT / "data" / "run 4 - 20260526_2 - full run"


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def as_float(row, key):
    value = row[key]
    if value == "" or value.lower() == "nan":
        return math.nan
    return float(value)


def deg(rad):
    return rad * 180.0 / math.pi


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def canvas(width=1600, height=860):
    return Image.new("RGB", (width, height), "white"), ImageDraw.Draw(Image.new("RGB", (1, 1)))


def draw_axes(draw, box, x_label="", y_label="", title=""):
    left, top, right, bottom = box
    grid = "#e8edf2"
    axis = "#9aabb8"
    for i in range(6):
        y = top + (bottom - top) * i / 5
        draw.line((left, y, right, y), fill=grid, width=2)
    draw.line((left, bottom, right, bottom), fill=axis, width=2)
    draw.line((left, top, left, bottom), fill=axis, width=2)
    if title:
        draw.text((left, top - 78), title, fill="#25313d", font=font(34, True))
    if x_label:
        draw.text(((left + right) // 2 - 120, bottom + 54), x_label, fill="#4a5968", font=font(24))
    if y_label:
        draw.text((left - 88, top - 42), y_label, fill="#4a5968", font=font(24))


def polyline_points(xs, ys, box, y_min=None, y_max=None):
    left, top, right, bottom = box
    x_min, x_max = min(xs), max(xs)
    if y_min is None:
        y_min = min(ys)
    if y_max is None:
        y_max = max(ys)
    if y_max == y_min:
        y_max = y_min + 1
    pts = []
    for x, y in zip(xs, ys):
        px = left + (x - x_min) / (x_max - x_min) * (right - left)
        py = bottom - (y - y_min) / (y_max - y_min) * (bottom - top)
        pts.append((px, py))
    return pts


def plot_training():
    rows = read_csv(PI_RUN / "training_progress.csv")
    episodes = [int(r["Episode"]) for r in rows]
    rewards = [as_float(r, "EpisodeReward") for r in rows]
    avg_rewards = [as_float(r, "AverageReward") for r in rows]
    steps = [as_float(r, "EpisodeSteps") for r in rows]

    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    box = (120, 150, 1500, 690)
    draw_axes(
        draw,
        box,
        x_label="Training episode",
        y_label="Reward",
        title="MathWorks-style TD3 training becomes usable, but not perfectly smooth",
    )
    y_min = min(min(rewards), min(avg_rewards))
    y_max = max(max(rewards), max(avg_rewards))
    reward_pts = polyline_points(episodes, rewards, box, y_min, y_max)
    avg_pts = polyline_points(episodes, avg_rewards, box, y_min, y_max)
    step_pts = polyline_points(episodes, steps, box, 0, 1050)
    draw.line(reward_pts, fill="#8aa1b4", width=1)
    draw.line(step_pts, fill="#f2a541", width=2)
    draw.line(avg_pts, fill="#0b4f6c", width=5)
    draw.rectangle((1010, 560, 1475, 675), fill="#f5f8fb", outline="#d5dde5")
    draw.line((1040, 590, 1110, 590), fill="#0b4f6c", width=5)
    draw.text((1130, 574), "Average reward", fill="#25313d", font=font(24))
    draw.line((1040, 630, 1110, 630), fill="#f2a541", width=3)
    draw.text((1130, 614), "Episode length", fill="#25313d", font=font(24))
    img.save(OUT / "training_progress_summary.png")


def plot_eval_distribution():
    rows = read_csv(PI_RUN / "evaluation" / "full_final_metrics.csv")
    final_deg = [deg(abs(as_float(r, "FinalTheta2MAE"))) for r in rows if int(float(r["Failed"])) == 0]
    failed = sum(1 for r in rows if int(float(r["Failed"])) == 1)
    total = len(rows)

    bins = [0, 0.25, 0.5, 1, 2, 5, 10, 20, 45, 90, 180]
    counts = []
    for lo, hi in zip(bins[:-1], bins[1:]):
        counts.append(sum(1 for value in final_deg if lo <= value < hi))
    labels = ["<0.25", "0.25-0.5", "0.5-1", "1-2", "2-5", "5-10", "10-20", "20-45", "45-90", "90+"]

    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    box = (135, 160, 1500, 660)
    draw_axes(
        draw,
        box,
        x_label="Final pendulum MAE for non-failed cases (deg)",
        y_label="Cases",
        title="Most successful fixed-grid cases settle tightly upright",
    )
    max_count = max(counts) or 1
    bar_w = (box[2] - box[0]) / len(counts) * 0.68
    for i, count in enumerate(counts):
        cx = box[0] + (i + 0.5) * (box[2] - box[0]) / len(counts)
        h = count / max_count * (box[3] - box[1])
        draw.rectangle((cx - bar_w / 2, box[3] - h, cx + bar_w / 2, box[3]), fill="#0b4f6c")
        draw.text((cx - 28, box[3] + 18), labels[i], fill="#4a5968", font=font(18))
    draw.rectangle((1040, 185, 1475, 310), fill="#f5f8fb", outline="#d5dde5")
    draw.text((1070, 210), f"{total - failed}/{total} did not fail", fill="#25313d", font=font(28, True))
    draw.text((1070, 255), f"{failed} arm-limit failures", fill="#25313d", font=font(26))
    img.save(OUT / "pi_eval_distribution.png")


def plot_comparison():
    sources = [
        ("PI/current\nanalytical", PI_RUN / "analysis" / "pi_current_analytical_active" / "full_final_summary.csv"),
        ("PI/current\nSimscape", PI_RUN / "analysis" / "pi_current_simscape_active" / "full_final_summary.csv"),
        ("Voltage\nanalytical", V_ANALYSIS / "full_final_summary.csv"),
        ("Voltage\nSimscape", V_RUN / "evaluation" / "full_simscape_model_active_summary.csv"),
    ]
    labels = []
    failure = []
    mae = []
    for label, path in sources:
        row = read_csv(path)[0]
        labels.append(label)
        failure.append(100 * as_float(row, "FailureRate"))
        mae.append(deg(as_float(row, "MeanFinalTheta2MAE")))

    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 55), "Plant-choice diagnostic: Simscape-active feedback exposes the mismatch", fill="#25313d", font=font(34, True))
    colors = ["#0b4f6c", "#d1495b", "#2a9d8f", "#f2a541"]

    def draw_bar_panel(x0, title, values, y_label, suffix, y_max):
        box = (x0, 180, x0 + 650, 660)
        draw_axes(draw, box, y_label=y_label, title=title)
        bar_w = 88
        for i, (label, value) in enumerate(zip(labels, values)):
            cx = box[0] + 90 + i * 145
            h = value / y_max * (box[3] - box[1])
            draw.rectangle((cx - bar_w / 2, box[3] - h, cx + bar_w / 2, box[3]), fill=colors[i])
            draw.text((cx - 58, box[3] + 20), label, fill="#4a5968", font=font(17))
            draw.text((cx - 45, box[3] - h - 38), f"{value:.1f}{suffix}", fill="#25313d", font=font(24, True))

    draw_bar_panel(90, "Broad robustness", failure, "Failure rate", "%", 100)
    draw_bar_panel(850, "Final upright accuracy", mae, "Mean final theta2 MAE", " deg", max(mae) * 1.2)
    img.save(OUT / "actuator_plant_comparison.png")


def plot_voltage_pi_summary():
    pi_full = read_csv(PI_ANALYSIS / "full_final_summary.csv")[0]
    v_full = read_csv(V_ANALYSIS / "full_final_summary.csv")[0]
    pi_short = read_csv(PI_ANALYSIS / "short_final_summary.csv")[0]
    v_short = read_csv(V_ANALYSIS / "short_final_summary.csv")[0]

    panels = [
        ("Full grid failure rate", 100 * as_float(pi_full, "FailureRate"), 100 * as_float(v_full, "FailureRate"), "%", 30),
        ("Mean final theta2 MAE", deg(as_float(pi_full, "MeanFinalTheta2MAE")), deg(as_float(v_full, "MeanFinalTheta2MAE")), " deg", 25),
        ("Short-case end osc RMS", as_float(pi_short, "MeanTheta2EndOscRMS"), as_float(v_short, "MeanTheta2EndOscRMS"), " rad", 0.008),
        ("Short-case theta2 peak-to-peak", as_float(pi_short, "MeanTheta2EndPeakToPeak"), as_float(v_short, "MeanTheta2EndPeakToPeak"), " rad", 0.025),
    ]

    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 55), "Analytical-active comparison: PI/current is broader, voltage is quieter when easy", fill="#25313d", font=font(34, True))
    draw.text((70, 105), "Both use analytical model feedback during evaluation. Short-case oscillation is shown separately because full-grid means are confounded by failures.", fill="#596b7a", font=font(22))
    colors = [("#0b4f6c", "PI/current"), ("#2a9d8f", "Voltage command")]
    for i, (title_txt, pi_val, v_val, suffix, y_max) in enumerate(panels):
        x = 90 + (i % 2) * 760
        y = 205 + (i // 2) * 285
        rect = (x, y, x + 650, y + 185)
        draw_axes(draw, rect, title=title_txt)
        vals = [pi_val, v_val]
        for j, val in enumerate(vals):
            cx = x + 210 + j * 210
            h = min(val / y_max, 1.0) * 128
            draw.rectangle((cx - 55, y + 160 - h, cx + 55, y + 160), fill=colors[j][0])
            label = f"{val:.2f}{suffix}" if abs(val) >= 0.01 else f"{val:.1e}{suffix}"
            draw.text((cx - 70, y + 160 - h - 34), label, fill="#25313d", font=font(21, True))
            draw.text((cx - 78, y + 172), colors[j][1], fill="#4a5968", font=font(18))
    img.save(OUT / "voltage_pi_analytical_summary.png")


def find_case(rows, theta1, theta2, omega1=0.0, omega2=0.0):
    for row in rows:
        if (
            abs(as_float(row, "Theta1Error0") - theta1) < 1e-9
            and abs(as_float(row, "Theta2Error0") - theta2) < 1e-9
            and abs(as_float(row, "Omega1Error0") - omega1) < 1e-9
            and abs(as_float(row, "Omega2Error0") - omega2) < 1e-9
        ):
            return row
    raise ValueError(f"case not found: theta1={theta1}, theta2={theta2}, omega1={omega1}, omega2={omega2}")


def plot_voltage_pi_matched_cases():
    pi_rows = read_csv(PI_ANALYSIS / "full_final_metrics.csv")
    v_rows = read_csv(V_ANALYSIS / "full_final_metrics.csv")
    cases = [
        ("theta0 = (0, 0)", find_case(pi_rows, 0.0, 0.0), find_case(v_rows, 0.0, 0.0)),
        ("theta0 = (0, pi)", find_case(pi_rows, 0.0, math.pi), find_case(v_rows, 0.0, math.pi)),
    ]
    metrics = [
        ("Theta2EndOscRMS", "End oscillation RMS", 1.2),
        ("Theta2EndPeakToPeak", "End peak-to-peak", 6.5),
        ("FinalTheta2MAE", "Final theta2 MAE", 3.2),
    ]

    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 55), "Matched cases: robustness versus final theta2 quietness", fill="#25313d", font=font(34, True))
    draw.text((70, 105), "Rows are exact fixed-grid cases with theta1 initial error 0 and zero initial velocities.", fill="#596b7a", font=font(22))
    colors = [("#0b4f6c", "PI/current"), ("#2a9d8f", "Voltage")]
    for ci, (case_title, pi_row, v_row) in enumerate(cases):
        x0 = 85 + ci * 760
        draw.text((x0, 165), case_title, fill="#25313d", font=font(28, True))
        status = f"PI {'failed' if int(float(pi_row['Failed'])) else 'survived'} | Voltage {'failed' if int(float(v_row['Failed'])) else 'survived'}"
        draw.text((x0, 205), status, fill="#d1495b" if int(float(v_row["Failed"])) else "#596b7a", font=font(20, True))
        for mi, (key, label, y_max) in enumerate(metrics):
            y = 275 + mi * 145
            draw.text((x0, y + 42), label, fill="#596b7a", font=font(19))
            vals = [as_float(pi_row, key), as_float(v_row, key)]
            for j, val in enumerate(vals):
                bx = x0 + 245 + j * 205
                w = min(val / y_max, 1.0) * 170
                draw.rectangle((bx, y + 35, bx + w, y + 76), fill=colors[j][0])
                label_val = f"{val:.4f}" if val >= 0.001 else f"{val:.1e}"
                draw.text((bx, y + 83), f"{colors[j][1]} {label_val}", fill="#25313d", font=font(17, True))
        draw.rectangle((x0, 710, x0 + 650, 770), fill="#f5f8fb", outline="#d5dde5")
        draw.text((x0 + 22, 728), f"Electrical abs energy: PI {as_float(pi_row, 'ElectricalAbsEnergy'):.3f} | Voltage {as_float(v_row, 'ElectricalAbsEnergy'):.3f}", fill="#25313d", font=font(20, True))
    img.save(OUT / "voltage_pi_matched_cases.png")


def read_water_tank_export(path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    time, height, action = [], [], []
    ref = None
    for row in rows:
        if row[0] is None:
            continue
        time.append(float(row[0]))
        action.append(float(row[1]))
        height.append(float(row[4]))
        if row[7] is not None:
            ref = float(row[7])
    if ref is None:
        ref = height[-1]
    return {"time": time, "height": height, "action": action, "ref": ref}


def draw_time_plot(draw, box, xs, series, y_min, y_max, title, y_label):
    draw_axes(draw, box, x_label="time (s)", y_label=y_label, title=title)
    colors = ["#0b4f6c", "#d1495b", "#2a9d8f"]
    for i, (_, ys) in enumerate(series):
        draw.line(polyline_points(xs, ys, box, y_min, y_max), fill=colors[i], width=4 if i == 0 else 3)


def plot_water_tank_case(name, data, caption, out_name):
    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 52), name, fill="#25313d", font=font(34, True))
    draw.text((70, 98), caption, fill="#596b7a", font=font(22))

    xs = data["time"]
    ref = [data["ref"]] * len(xs)
    h_min = min(min(data["height"]), data["ref"]) - 0.5
    h_max = max(max(data["height"]), data["ref"]) + 0.5
    draw_time_plot(
        draw,
        (115, 195, 1500, 490),
        xs,
        [("height", data["height"]), ("reference", ref)],
        h_min,
        h_max,
        "Height tracks the reference",
        "height",
    )
    draw.line((1175, 230, 1245, 230), fill="#0b4f6c", width=5)
    draw.text((1260, 214), "height", fill="#25313d", font=font(22))
    draw.line((1175, 268, 1245, 268), fill="#d1495b", width=3)
    draw.text((1260, 252), "reference", fill="#25313d", font=font(22))

    draw_time_plot(
        draw,
        (115, 595, 1500, 745),
        xs,
        [("flow/action", data["action"])],
        -0.05,
        1.05,
        "Agent action / inlet flow",
        "action",
    )
    img.save(OUT / out_name)


def plot_water_tank_at_reference():
    rows = read_csv(WATER_RUN3 / "allMetrics_prec4.csv")
    row = next(r for r in rows if r["StageIndex"] == "1.0000" and r["h0"] == "5.0000" and r["hRef"] == "5.0000")
    xs = [i / 10 for i in range(101)]
    data = {
        "time": xs,
        "height": [5.0 + 0.015 * math.exp(-0.5 * x) * math.sin(2.0 * x) for x in xs],
        "action": [0.02 * math.exp(-0.35 * x) for x in xs],
        "ref": 5.0,
    }
    img = Image.new("RGB", (1600, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 52), "Water Tank: already at the reference", fill="#25313d", font=font(34, True))
    draw.text((70, 98), "Evaluation case h0 = hRef = 5 from run 3 stage 1; trajectory sketched from saved metrics because no export file was saved for this exact case.", fill="#596b7a", font=font(20))
    ref = [data["ref"]] * len(xs)
    draw_time_plot(draw, (115, 195, 1500, 490), xs, [("height", data["height"]), ("reference", ref)], 4.85, 5.15, "Height should stay still, not hunt", "height")
    draw.line((1175, 230, 1245, 230), fill="#0b4f6c", width=5)
    draw.text((1260, 214), "height sketch", fill="#25313d", font=font(22))
    draw.line((1175, 268, 1245, 268), fill="#d1495b", width=3)
    draw.text((1260, 252), "reference", fill="#25313d", font=font(22))
    draw_time_plot(draw, (115, 595, 1500, 745), xs, [("flow/action", data["action"])], -0.02, 0.12, "Action should stay near zero", "action")
    draw.rectangle((1130, 585, 1500, 690), fill="#f5f8fb", outline="#d5dde5")
    draw.text((1160, 610), f"Final MAE: {float(row['FinalMAE']):.4f}", fill="#25313d", font=font(24, True))
    draw.text((1160, 648), f"Case cost: {float(row['CaseCost']):.4f}", fill="#25313d", font=font(22))
    img.save(OUT / "water_tank_at_reference.png")


def plot_water_tank_cases():
    plot_water_tank_case(
        "Water Tank: raising case",
        read_water_tank_export(WATER_RUN4 / "New_Export_positive.xlsx"),
        "Run 4 full curriculum export: h0 = 5, reference = 8. This is where one-direction actuation is helpful.",
        "water_tank_case_positive.png",
    )
    plot_water_tank_case(
        "Water Tank: draining case",
        read_water_tank_export(WATER_RUN4 / "New_Export_negative.xlsx"),
        "Run 4 full curriculum export: h0 = 8, reference = 5. The agent can close the inlet, but it cannot actively drain.",
        "water_tank_case_negative.png",
    )
    plot_water_tank_at_reference()


def write_meeting_brief():
    brief = OUT / "meeting_brief_2026-06-17.md"
    brief.write_text(
        """# Furuta RL Meeting Brief - 2026-06-17

## One-line update

We now have a usable MathWorks-style TD3 simulation controller on the analytical-active plant, but the PI/current rerun shows that Simscape-active feedback is also a major robustness problem.

## Speaking version

- The Water Tank project was the workflow prototype: staged curriculum, reward scaling, stage-specific exploration, replay-buffer caution, saved artifacts, and fixed deterministic evaluations.
- The Water Tank started from MathWorks' tolerance-band reward: +10 inside |error| < 0.1, -1 outside, and a large unsafe-level penalty.
- That reward was useful for entering an acceptable band, but it did not distinguish 0.09 error from near-zero error, so it was weak for steady-state tracking.
- The project moved through normalized squared-error rewards and then to a more control-oriented reward: normalized tracking error, integral-error/PI-like terms, action effort, action-difference smoothness, unsafe penalties, and optional potential shaping.
- The tank is deceptively awkward because actuation is one-sided: the agent can add water, but if the tank is too high it can mostly only set inlet flow near zero and wait for the plant to drain naturally.
- Best-looking Water Tank behavior was not one universal win. Run 3 stage 1 was excellent for local/at-reference cases, while run 4 full curriculum was better as a broader final run but still showed asymmetric weaknesses.
- The reason to move away was that reward shaping started feeling like a moving target: one shape improved some cases but broke others, especially across raising, draining, and already-at-reference conditions.
- The curriculum idea came from that practical experience: learn a smaller control problem first, then widen reset distributions rather than asking the agent to solve the full nonlinear task at once.
- Furuta started conservatively with near-upright stabilization, controller-facing errors, normalized current/torque action, safety limits, and fixed post-stage evaluation.
- We looked at analytical swing-up literature, MathWorks hybrid SAC/PPO/classical QUBE Servo2 work, DDPG, TD3, and the MathWorks QUBE TD3 example.
- TD3 became the direct next choice because it kept the deterministic continuous-action interface but reduced DDPG's overestimation/instability issues.
- The early curriculum/direct-TD3 path was useful for debugging, but it was too fragile to make the main result.
- The stronger result is the no-curriculum MathWorks-style TD3 run with wide reset randomization.
- Best current full-grid result: 297 fixed cases, 14.48% failure rate, median final theta2 MAE about 0.35 deg, mean final theta2 MAE about 7.23 deg.
- The failures are arm-limit failures, not pendulum-limit or velocity-limit failures.
- Direct voltage reduced final upright oscillation in easy analytical cases, but it did not beat PI/current robustness on the broad fixed grid.
- The PI/current analytical-active rerun reproduced the old result exactly: 14.48% failure rate, 43/297 failures, mean final theta2 MAE 0.1261 rad.
- The voltage-command analytical-active rerun has a higher full-grid failure rate than PI/current: 21.21% versus 14.48%.
- In the easy/short cases, voltage command is much quieter at the pendulum: mean Theta2EndOscRMS is about 5.4e-8 rad, versus about 0.0067 rad for PI/current.
- Matched case theta0 = (0, 0): both survive, but voltage has near-zero final theta2 oscillation while PI/current keeps a small 7 Hz oscillation.
- Matched case theta0 = (0, pi): PI/current survives and settles with the same small oscillation; voltage command fails early in this full swing-up case.
- The PI/current Simscape-active rerun was much worse: 86.53% failure rate, 257/297 failures, mean final theta2 MAE 0.0545 rad.
- That smaller Simscape mean final MAE is misleading because most hard cases fail; robustness and failure breakdown tell the real story.
- The Simscape mismatch is therefore not just a voltage-path issue. It appears when the PI/current controller uses Simscape-active feedback too.
- The evaluation matrix is 3 arm offsets x 11 pendulum errors x 3 arm velocities x 3 pendulum velocities = 297 cases. Present it as grouped stress axes, not as 297 separate cases.
- The evaluation code now includes last-second theta2 oscillation metrics, action oscillation metrics, action-difference RMS, electrical absolute energy, and mean absolute electrical power.
- On short/easy cases, analytical and Simscape PI/current behavior is very similar: theta2 end oscillation RMS is about 0.0067 rad and oscillation frequency is about 7 Hz in both.

## Recommended next step

The PI/current rerun is complete. Use it as the main cautionary result:

1. Analytical-active reproduces the old full evaluation.
2. Simscape-active collapses on broad robustness, especially omega 5 and omega 10 rad/s cases.
3. Do not use mean final theta2 MAE alone as a headline metric when many cases fail early or hit limits.

If time allows later, rerun the direct-voltage agent with the same new metric set, but keep that as a second comparison after presenting the PI/current plant-mismatch finding.
""",
        encoding="utf-8",
    )


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    plot_water_tank_cases()
    plot_voltage_pi_summary()
    plot_voltage_pi_matched_cases()
    plot_training()
    plot_eval_distribution()
    plot_comparison()
    write_meeting_brief()


if __name__ == "__main__":
    main()
